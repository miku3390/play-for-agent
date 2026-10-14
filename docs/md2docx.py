from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()

# 设置默认中文字体
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)

# 读取 Markdown 文件
with open('D:\\Desktop\\CODEX-d4\\docs\\handkerchief-arena-architecture.md', 'r', encoding='utf-8') as f:
    content = f.read()

lines = content.split('\n')

i = 0
in_code_block = False
code_lines = []

def add_code_block():
    global code_lines
    if code_lines:
        code_text = '\n'.join(code_lines)
        p = doc.add_paragraph()
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        code_lines = []

while i < len(lines):
    line = lines[i]
    
    # 代码块
    if line.strip().startswith('```'):
        if in_code_block:
            add_code_block()
            in_code_block = False
        else:
            in_code_block = True
        i += 1
        continue
    
    if in_code_block:
        code_lines.append(line)
        i += 1
        continue
    
    # 空行
    if line.strip() == '':
        i += 1
        continue
    
    # 标题
    heading_match = re.match(r'^(#{1,4})\s+(.+)$', line)
    if heading_match:
        level = len(heading_match.group(1))
        text = heading_match.group(2)
        if level == 1:
            doc.add_heading(text, level=0)
        else:
            doc.add_heading(text, level=level)
        i += 1
        continue
    
    # 表格
    if '|' in line and i + 1 < len(lines) and re.match(r'\|[-:\s]+\|', lines[i+1]):
        table_lines = [line]
        i += 1
        # 跳过分隔行
        if i < len(lines) and re.match(r'\|[-:\s]+\|', lines[i]):
            i += 1
        while i < len(lines) and '|' in lines[i]:
            table_lines.append(lines[i])
            i += 1
        
        rows = []
        for tl in table_lines:
            cells = [c.strip() for c in tl.split('|')[1:-1]]
            if cells:
                rows.append(cells)
        
        if rows:
            num_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = 'Light Grid Accent 1'
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        table.rows[ri].cells[ci].text = cell_text
        continue
    
    # 普通段落（支持加粗 **text**）
    p = doc.add_paragraph()
    parts = re.split(r'(\*\*.*?\*\*)', line)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)
    
    i += 1

# 保存
doc.save('D:\\Desktop\\CODEX-d4\\docs\\handkerchief-arena-architecture.docx')
print('Word document created successfully')
